"""Avviksliste: one workbook per round, one sheet per requirement, only the flagged elements.

    python bygg_avvik.py --prosjekt krav.yaml --runde runde.json --cache CACHE --ut UT
        -> <ut>/<kode>_Mottakskontroll_Avvik_<fag>.xlsx, one per discipline group,
           beside the main workbook

edkjo 2026-09-25: by DISCIPLINE, not across disciplines. The groups come from krav.yaml
`faggrupper`, an explicit list group -> model labels (config, never derived from the name); a model
listed nowhere is its own group, one file per model. The MMI subsets sit in their
discipline. Each file has one sheet per requirement; the Modell column says which subset a row
comes from. A flagged item is an element that lacks the property (against Dekning) or carries
an invalid value (against Gyldig). Columns: Modell · Fagkode · GUID · IFC-klasse · Navn · Typenavn ·
Etasje · Egenskap · Funnet verdi · Årsak. Sheets follow the report's numbering.

A data export, not IFC processing: it reads only what is cached from the round,
    <ut>/data/beregnet.json                    requirements, their order and locations, etasjematrise
    <cache>/bep/<dato>/data.json               every counted GlobalId with its HI90_Kopi objekt value
    <cache>/blokker/<dato>/prosjekt/<modell>.json  per-element findings of each requirement at its one
                                               location (blokkdata 6), with class, name, type, storey
Where a requirement has no per-element list in these caches, its sheet says so in one visible row;
nothing is measured to fill it. Requirements the report draws as components (GUID, Etasjedefinisjon)
keep their sheet, unnumbered, as in the report.
"""
from __future__ import annotations

import argparse
import collections
import json
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
HER = Path(__file__).resolve().parent
sys.path.insert(0, str(HER))
import bygg_mottakskontroll as B  # noqa: E402
import stier  # noqa: E402

KOLONNE_NOKLER = ("modell", "fagkode", "guid", "ifc_klasse", "navn", "typenavn", "etasje", "egenskap",
                  "funnet_verdi", "aarsak")   # labels from `etiketter`


def faggruppe(label: str, fg: dict[str, list[str]]) -> str:
    """A model's discipline group, looked up in krav.yaml `faggrupper` (group -> labels). A model
    listed nowhere is its own group (edkjo: «Unless explicitly configured, see every file as
    separate»); nothing is derived from the name."""
    treff = [g for g, labels in fg.items() if label in (labels or [])]
    if len(treff) > 1:
        raise SystemExit(f"FEIL: {label} står i flere faggrupper: {', '.join(treff)}")
    return treff[0] if treff else label


def bygg(prosjekt: Path, runde: Path, cache: Path, ut: Path, kun_standard: bool = False) -> int:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    stier.finnes(prosjekt, runde)
    reg = B.Register(prosjekt, kun_standard)
    r = B.Runde(runde, cache, prosjekt.resolve().parent)
    r.bind(reg)
    R = json.loads((ut / "data" / "beregnet.json").read_text(encoding="utf-8"))
    data = {d["label"]: d for d in json.loads(r.data_sti.read_text(encoding="utf-8"))}
    blokk = {}
    for m in R["modeller"]:
        p = r.blokk_sti / ("standard" if kun_standard else "prosjekt") / f"{m['label']}.json"
        if p.exists():
            d_ = json.loads(p.read_text(encoding="utf-8"))
            if d_.get("ett_sted"):
                blokk[m["label"]] = d_
    AARSAK = {"mangler": ("avvik_mangler", "bad"), "tom": ("tom_verdi", "warn"),
              "ugyldig": ("ugyldig_verdi", "warn"), "ikke_i_etasje": ("ikke_i_etasje", "warn")}

    fg = reg.faggrupper   # from the merged config; empty in a standard-only run
    if not isinstance(fg, dict) or any(not isinstance(v, list) for v in fg.values()):
        raise SystemExit("FEIL: faggrupper i krav.yaml skal være gruppe -> [modellnavn]")
    grupper: dict[str, list[str]] = {}
    for m in R["modeller"]:
        grupper.setdefault(faggruppe(m["label"], fg), []).append(m["label"])
    tok = B.les_tokens()
    fyll = {s: PatternFill("solid", fgColor=B.hex_(tok, f"--sk-{n}-flate"))
            for s, n in (("bad", "mangler"), ("warn", "avvik"), ("na", "na"))}
    bdefs = {b["id"]: b for b in reg.blokker}
    labels = [m["label"] for m in R["modeller"]]
    alle_guid = B.guid_forekomster(data, labels)
    rapport = []

    def objekt(lab: str, guid: str) -> list[str]:
        """[IFC class, name, type name, storey] of a listed element, from the measurement."""
        return (blokk.get(lab) or {}).get("objekter", {}).get(guid) or ["", "", "", ""]

    def etasje_for(lab: str, guid: str) -> str:
        return objekt(lab, guid)[3]

    ark_liste: list[tuple[str, str, list[list], bool]] = []

    def ark(bid: str, tittel: str, rader: list[list], hull: bool = False):
        ark_liste.append((bid, tittel, rader, hull))

    def skriv_ark(wb, tittel: str, rader: list[list], hull: bool):
        ws = wb.create_sheet(tittel[:31])
        ws.append([B.E(k) for k in KOLONNE_NOKLER])
        for c in ws[1]:
            c.font = Font(bold=True)
            c.alignment = Alignment(wrap_text=True, vertical="bottom")
        if hull:
            ws.append(["", "", "", "", "", "", "", "", "", B.E("hull")])
            ws.cell(row=2, column=10).fill = fyll["na"]
        rad_nr = ws.max_row
        for rad in rader:
            ws.append(rad[:10])
            rad_nr += 1   # counted here: ws.max_row scans every cell, quadratic on long lists
            ws.cell(row=rad_nr, column=10).fill = fyll[rad[10]]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(KOLONNE_NOKLER))}{max(ws.max_row, 1)}"
        for i, w in enumerate((20, 8, 24, 20, 24, 24, 14, 32, 24, 30), start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        return 0 if hull else len(rader)

    nr = 0
    komp = set(R["modeller"][0].get("komponent_ids") or [])
    for bid in [b["id"] for b in R["modeller"][0]["blokker"]]:
        bdef = bdefs[bid]
        if bid in komp:
            tittel = bdef["tittel"]
        else:
            nr += 1
            tittel = f"{nr} {bdef['tittel']}"
        sted = B.blokkdata.kildeetikett(bdef["sted"]) if bdef.get("sted") else ""
        rader: list[list] = []
        hull = False
        if bdef.get("ikke_konfigurert"):
            # no location in either layer (e.g. MMI, Kopiobjekt in a standard-only run): one row
            rader = [["alle", "", "", "", "", "", "", "", "", B.E("ikke_konfigurert"), "na"]]
        elif bid == "guid":
            # duplicates inside a file (none in this round), then the GUIDs shared with other models
            for m in R["modeller"]:
                lab = m["label"]
                if int(data[lab].get("dupliserte_guid") or 0):
                    hull = True   # the cache keeps the count of duplicates, not the list
                for g in data[lab].get("guid_kopi") or {}:
                    andre = [x for x, _ in alle_guid.get(g, []) if x != lab]
                    if andre:
                        rader.append([lab, m["fag"], g, "", "", "", etasje_for(lab, g), sted,
                                      data[lab]["guid_kopi"][g] or "",
                                      B.E("delt_med") + " " + ", ".join(B.kort_label(x, reg) for x in andre), "warn"])
        elif bid == "etasjedefinisjon":
            for m in R["modeller"]:
                for e in m["etasjematrise"]["rader"]:
                    if e["kode"] == "som_registeret":
                        continue
                    funnet = f'{e["m_navn"]} {B.m_(e["m_kote"], 3)}' if e["m_navn"] else ""
                    ref = f'{e["k_navn"]} {B.m_(e["k_kote"], 3)}' if e["k_navn"] else ""
                    st = "na" if e["kode"] == "mangler" else "bad"
                    rader.append([m["label"], m["fag"], "", "IfcBuildingStorey", e["m_navn"], "", ref,
                                  "IfcBuildingStorey", funnet, e["status"], st])
        elif bid in ("typeobjekt", "etasjer", "systemkode", "funksjonskode", "produkt", "materiale", "mmi", "fase"):
            # per-element findings at the requirement's one location (blokkdata 6)
            for m in R["modeller"]:
                if next((b["verdikt"] for b in m["blokker"] if b["id"] == bid), "") == "gjelder_ikke":
                    continue   # not applicable to this model: nothing in it is a deviation
                d_ = blokk.get(m["label"])
                if d_ is None:
                    hull = True
                    continue
                for guid, verdi, kode in (d_.get("funn") or {}).get(bid, []):
                    kl, navn, tn, et = objekt(m["label"], guid)
                    etikett, st = AARSAK[kode]
                    rader.append([m["label"], m["fag"], guid, kl, navn, tn, et, sted, verdi, B.E(etikett), st])
        elif bid == "kopiobjekt":
            godtatte = {f.upper() for f in reg.fagkoder}
            sanne = {str(x).lower() for x in reg.forventet["kopi_objekt"]["sanne_verdier"]}
            for m in R["modeller"]:
                d = data[m["label"]]
                finnes = int((d.get("finnes") or {}).get("HI90_Kopi objekt", 0))
                for guid, verdi in (d.get("guid_kopi") or {}).items():
                    v = (verdi or "").strip()
                    if not v:
                        # the cache stores absent and empty alike; the per-model count tells which
                        arsak = ("avvik_mangler" if finnes == 0 else "tom_verdi" if finnes >= d["n"]
                                 else "mangler_eller_tom_verdi")
                        rader.append([m["label"], m["fag"], guid, "", "", "", etasje_for(m["label"], guid), sted, "",
                                      B.E(arsak), "bad" if arsak == "avvik_mangler" else "warn"])
                    elif v.upper() not in godtatte and v.lower() not in sanne:
                        rader.append([m["label"], m["fag"], guid, "", "", "", etasje_for(m["label"], guid), sted, v,
                                      B.E("ugyldig_verdi"), "warn"])
        else:
            hull = True
        ark(bid, tittel, rader, hull)

    def slaa_sammen(rader: list[list]) -> list[list]:
        """Within a discipline an identical finding is listed once, with every subset that carries it
        in the Modell column. The key is the finding's content (all columns but Modell), never the
        model name; per-element findings stay apart because their GUIDs differ."""
        samlet: dict[tuple, list] = {}
        for r_ in rader:
            nokkel = tuple(r_[1:])
            if nokkel in samlet:
                samlet[nokkel][0].append(r_[0])
            else:
                samlet[nokkel] = [[r_[0]], r_]
        return [[" · ".join(B.kort_label(x, reg) for x in modeller)] + r_[1:] for modeller, r_ in samlet.values()]

    def guid_rader(medlemmer: list[str]) -> list[list]:
        """Shared GUIDs of a discipline: one row per GUID with the subsets that carry it, and the
        models outside the discipline it is shared with."""
        per_guid: dict[str, list[str]] = {}
        for lab in medlemmer:
            for g in data[lab].get("guid_kopi") or {}:
                if len(alle_guid.get(g, [])) > 1:
                    per_guid.setdefault(g, []).append(lab)
        ut_ = []
        for g, egne in per_guid.items():
            andre = [x for x, _ in alle_guid[g] if x not in medlemmer]
            arsak = ((B.E("delt_med") + " " + ", ".join(B.kort_label(x, reg) for x in andre)) if andre
                     else B.E("delt_mellom_delmodellene"))
            lab = egne[0]
            ut_.append([" · ".join(B.kort_label(x, reg) for x in egne),
                        next(m["fag"] for m in R["modeller"] if m["label"] == lab), g, "", "", "",
                        etasje_for(lab, g), B.blokkdata.kildeetikett(bdefs["guid"]["sted"]),
                        data[lab]["guid_kopi"][g] or "", arsak, "warn"])
        return ut_

    gammel = ut / f"{reg.prosjekt['kode']}_Mottakskontroll_Avvik.xlsx"
    feil = []
    if gammel.exists():
        try:
            gammel.unlink()   # the combined file is replaced by one per discipline
        except PermissionError:
            feil.append(f"den samlede fila er åpen og ble ikke slettet: {gammel}")
    skrevet = []
    for fag, medlemmer in grupper.items():
        wb = openpyxl.Workbook()
        wb.remove(wb["Sheet"])
        tall = []
        for bid, tittel, rader, hull in ark_liste:
            if bid == "guid":
                egne = guid_rader(medlemmer)
            elif rader and rader[0][0] == "alle":
                egne = rader
            else:
                egne = slaa_sammen([r_ for r_ in rader if r_[0] in medlemmer])
            tall.append((tittel, skriv_ark(wb, tittel, egne, hull), hull))
        sti = ut / f"{reg.prosjekt['kode']}_Mottakskontroll_Avvik_{fag}.xlsx"
        try:
            wb.save(sti)
        except PermissionError:
            sti = ut / "html" / f"{reg.prosjekt['kode']}_Mottakskontroll_Avvik_{fag}_1.xlsx"
            wb.save(sti)
            feil.append(f"arbeidsboka for {fag} er låst (åpen); ny versjon ligger i {sti}")
        skrevet.append(sti)
        print(f"{sti}  ({', '.join(medlemmer)})")
        for navn, n, hull in tall:
            print(f"  {navn}: {'HULL' if hull else n}")
    feil += [f"mojibake: {x}" for x in B.sjekk_mojibake(skrevet)]
    for x in feil:
        print("FEIL:", x)
    return 1 if feil else 0


def main() -> int:
    ap = argparse.ArgumentParser()
    stier.argumenter(ap)
    ap.add_argument("--kun-standard", action="store_true",
                    help="the standard layer only; reads <ut>/data/beregnet.json of a --kun-standard build")
    a = ap.parse_args()
    return bygg(a.prosjekt, a.runde, a.cache, a.ut, a.kun_standard)


if __name__ == "__main__":
    raise SystemExit(main())
