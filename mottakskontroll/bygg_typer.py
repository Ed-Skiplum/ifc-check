"""The type report: the mottakskontroll per type, one row per type, as a workbook.

    python bygg_typer.py --prosjekt krav.yaml --runde runde.json --cache CACHE --ut DIR [--ifc DIR]

One data source, two views (edkjo): the model report (bygg_mottakskontroll.py) aggregates the
per-object readings blokkdata.mål makes; this report groups the same readings by type. Nothing is
parsed here. mål writes every reading of every object to the cache (`per_objekt`, only for a project
with `typerapport`); each type's readings are replayed into a Blokk of its own and judged by the
model report's own code, Modellberegning.krav, with the same thresholds, `unntatt`,
`gjelder_fra_mmi`, avvik/mangler split and validity. The build stops unless, per model and
requirement, the per-type counts sum to the model report's (<ut>/data/beregnet.json): oppfylt,
avvik, mangler, gjelder ikke, and the objects inside and outside the accepted values.

A type is the model + IfcTypeObject (IFC class + name, begreper §1); the objects without a type are
one row per model and IFC class. The rows are the types with counted objects (the counted set,
standard.yaml `standard`). Parts of an assembly (`standard.ikke_telt_del_av`) are not counted and
not on the type rows; they are listed among the objects with their owner.

Sheets: Sammendrag (per model), Typer (per type: the requirements as columns, grouped by the
project's report sections), Forekomster (per object: its values and status per requirement, and the
parts with their owner), Materialer, Materiallag, Kilder. Labels from standard.yaml `etiketter`
(`ty_*`, English through standard_en.yaml).

Project config, key `typerapport` (its presence turns the report on):
    filnavn            workbook name (default <kode>_Mottakskontroll_Typer.xlsx)
    seksjoner          report section ids whose requirements get columns (default: every section)
    utelatt_modeller   {label: reason}: models with no rows (listed in Kilder)
    mengder            [{id, tittel, navn: [quantity names], enhet: m3 | m2 | m}]: the objects' own
                       base quantities (the first listed name a quantity set carries), summed per type
    klar               {tittel, blokker: [requirement ids]}: a type is «klar» when no object of it is avvik or
                       mangler and no value is outside the accepted values in any of them
"""
from __future__ import annotations

import argparse
import collections
import datetime as dt
import hashlib
import json
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
HER = Path(__file__).resolve().parent
sys.path.insert(0, str(HER))
import blokkdata  # noqa: E402
import bygg_mottakskontroll as MK  # noqa: E402
import konfig  # noqa: E402
import stier  # noqa: E402

BOTTER = ("oppfylt", "avvik", "mangler", "gjelder_ikke")
# Model-level requirements: no per-object reading, so no column and no per-type count.
IKKE_PER_OBJEKT = ("guid", "etasjedefinisjon")
NOKLER = ("filnavn", "seksjoner", "utelatt_modeller", "mengder", "klar")
E = MK.E


def oppsett(y: dict, reg: MK.Register) -> dict:
    """`typerapport` checked: unknown keys, sections, requirement ids and units stop the run."""
    tr = y.get("typerapport") or {}
    ukjent = set(tr) - set(NOKLER)
    if ukjent:
        raise SystemExit(f"FEIL: typerapport: ukjente nøkler {sorted(ukjent)} (gyldige: {', '.join(NOKLER)})")
    sids = [s["id"] for s in reg.seksjoner]
    valgt = [str(s) for s in tr.get("seksjoner") or sids]
    if set(valgt) - set(sids):
        raise SystemExit(f"FEIL: typerapport.seksjoner: {sorted(set(valgt) - set(sids))} finnes ikke ({sids})")
    kol = [b for b in reg.blokker if b["id"] not in IKKE_PER_OBJEKT + ("typeobjekt",)
           and b.get("seksjon") in valgt]
    kol.sort(key=lambda b: valgt.index(b["seksjon"]))
    klar = [str(x) for x in (tr.get("klar") or {}).get("blokker") or []]
    if set(klar) - {b["id"] for b in kol}:
        raise SystemExit(f"FEIL: typerapport.klar.blokker: {sorted(set(klar) - {b['id'] for b in kol})} "
                         "er ikke krav med kolonne i typerapporten")
    for m in tr.get("mengder") or []:
        if not m.get("id") or not m.get("tittel") or not m.get("navn") or m.get("enhet") not in blokkdata.ENHETSTYPE:
            raise SystemExit(f"FEIL: typerapport.mengder: {m!r} trenger id, tittel, navn og enhet "
                             f"({', '.join(blokkdata.ENHETSTYPE)})")
    return {"filnavn": tr.get("filnavn") or f"{reg.prosjekt['kode']}_Mottakskontroll_Typer.xlsx",
            "kolonner": kol, "seksjoner": [s for s in reg.seksjoner if s["id"] in valgt],
            "utelatt": {str(k): str(v) for k, v in (tr.get("utelatt_modeller") or {}).items()},
            "mengder": tr.get("mengder") or [], "klar": klar, "klar_paa": "klar" in tr,
            "klar_tittel": str((tr.get("klar") or {}).get("tittel") or E("ty_klar_kol"))}


# ================================================================ per type, by the model report's code
def gruppenokkel(o: dict) -> tuple[str, str]:
    return (o["t"][0], o["t"][1]) if o["t"] else ("", o["k"])


def spill(objekter: list[dict], blokkinfo: dict) -> dict[str, blokkdata.Blokk]:
    """The objects' readings replayed, in mål's order, into fresh Blokk objects."""
    blk = {bid: blokkdata.Blokk(list(v["kildenavn"]), v["ikke_konfigurert"]) for bid, v in blokkinfo.items()}
    for o in objekter:
        tk = tuple(o["t"]) if o["t"] else None
        for ev in o["b"]:
            bid, art = ev[0], ev[1]
            if art == "s":
                blk[bid].sett(ev[2], ev[3], ev[4], tk, ev[5])
            elif art == "e":
                blk[bid].ekstra[ev[2]] += 1
            else:
                blk[bid].mmi_utenfor[ev[2]] += 1
    return blk


def gruppe_bd(objekter: list[dict], blokkinfo: dict, typet: tuple[str, str] | None) -> tuple[dict, dict]:
    """A blokkdata-shaped dict for a set of objects: what Modellberegning.krav reads."""
    blk = spill(objekter, blokkinfo)
    telte = [o for o in objekter if o["c"]]
    mk = collections.Counter(tuple(o["mk"]) for o in telte)
    et = collections.Counter(tuple(o["s"]) for o in telte if o.get("s"))
    bd = {"n": len(telte), "n_alle": len(objekter), "ett_sted": True,
          "blokker": {bid: b.ut() for bid, b in blk.items()},
          "mengdeklasser": [[kl, mt, en, c, ap] for (kl, mt, en, ap), c in mk.items()],
          "etasjer": [{"navn": n_, "elevation_mm": el, "objekter": c} for (n_, el), c in et.items()],
          "hierarki": dict(collections.Counter(o.get("h") for o in telte)),
          "typer_alle": [[typet[0], typet[1], len(objekter)]] if typet else [],
          "typeregister": {"liste": [], "typer": 0, "typede": 0}, "per_klasse": {}, "type_1til1": {},
          "materialer": []}
    return bd, blk


def dommer(mb: MK.Modellberegning, d0: dict, bd: dict, bdefs: list[dict]) -> dict[str, dict]:
    mb.bd, mb.d = bd, {**d0, "n": bd["n"]}
    return {b["id"]: mb.krav(b) for b in bdefs}


def tall(b: dict) -> dict:
    """The counts parity compares: the buckets, and the objects inside / outside the accepted values."""
    ut = {k: int(b["dekning"].get(k, 0)) for k in BOTTER}
    gy = (b.get("rader") or [{}])[0].get("gyldighet") if b.get("rader") else None
    if gy is not None:
        ut["innenfor"] = int(gy["innenfor"])
        ut["utenfor"] = sum(int(n) for _, n in gy["utenfor"])
    return ut


# ================================================================ the build
def bygg(prosjekt: Path, runde: Path, cache: Path, ut: Path, ifc: Path | None = None) -> int:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    stier.finnes(prosjekt, runde)
    reg = MK.Register(prosjekt)
    y = konfig.last(prosjekt)
    if not y.get("typerapport"):
        print("typerapport er ikke konfigurert: ingen typerapport")
        return 0
    T = oppsett(y, reg)
    r = MK.Runde(runde, cache, prosjekt.resolve().parent, ifc)
    r.bind(reg)
    beregnet = ut / "data" / "beregnet.json"
    stier.finnes(r.data_sti, r.struktur_sti, beregnet)
    R = json.loads(beregnet.read_text(encoding="utf-8"))
    rm = {m["label"]: m for m in R["modeller"]}
    data = {d["id"]: d for d in map(MK.normaliser, json.loads(r.data_sti.read_text(encoding="utf-8")))}
    struktur = {s["label"]: s for s in json.loads(r.struktur_sti.read_text(encoding="utf-8"))}
    hoved = r.hoved(reg)
    labels = [lab for lab in hoved if lab not in T["utelatt"]]
    tok = MK.les_tokens()
    kodetabeller = {b["id"]: {str(x["kode"]).strip().upper() for x in __import__("yaml").safe_load(
                        stier.tabell_sti(b["kodetabell"], reg.rot).read_text(encoding="utf-8"))}
                    for b in reg.blokker if b.get("kodetabell")}
    alle_guid = MK.guid_forekomster(data, hoved)
    paritet = [b for b in reg.blokker if b["id"] not in IKKE_PER_OBJEKT]
    kol = T["kolonner"]
    print(f"Typerapport {r.dato}: {len(labels)} modeller, {len(kol)} krav som kolonner")

    rader, forekomster, feil, sammendrag, deler_alle = [], [], [], [], []
    for lab in labels:
        bd0 = json.loads((r.blokk_sti / "prosjekt" / f"{lab}.json").read_text(encoding="utf-8"))
        po = bd0.get("per_objekt")
        if not po or po.get("versjon") != blokkdata.PER_OBJEKT_VERSJON:
            raise SystemExit(f"FEIL: {lab}: målingen har ingen logg per objekt (versjon "
                             f"{blokkdata.PER_OBJEKT_VERSJON}); mål runden på nytt med typerapport konfigurert")
        if bd0["n"] != rm[lab]["n"]:
            raise SystemExit(f"FEIL: {lab}: målingen n={bd0['n']} ≠ modellrapportens n={rm[lab]['n']}")
        blokkinfo = {bid: {"kildenavn": v["kildenavn"], "ikke_konfigurert": v["ikke_konfigurert"]}
                     for bid, v in bd0["blokker"].items()}
        mb = MK.Modellberegning(reg, r, lab, data[lab], bd0, struktur[lab], alle_guid, tok, kodetabeller, {})
        grupper: dict[tuple[str, str], list[dict]] = collections.defaultdict(list)
        for o in po["objekter"]:
            grupper[gruppenokkel(o)].append(o)
        # per type: the model report's judgement, and the parity sum
        sum_ = {b["id"]: collections.Counter() for b in paritet}
        mine = []
        for (tkl, tnv), objs in grupper.items():
            bd, blk = gruppe_bd(objs, blokkinfo, (tkl, tnv) if tkl else None)
            dom = dommer(mb, data[lab], bd, paritet)
            for bid, b in dom.items():
                sum_[bid].update(tall(b))
            if bd["n"]:
                mine.append({"modell": lab, "typet": bool(tkl), "typeklasse": tkl, "typenavn": tnv if tkl else "",
                             "objekter": [o for o in objs if o["c"]], "n": bd["n"], "dom": dom, "blk": blk})
        for b in paritet:
            rb = next(x for x in rm[lab]["blokker"] if x["id"] == b["id"])
            vent = {k: v for k, v in tall(rb).items()}
            har = {k: sum_[b["id"]].get(k, 0) for k in vent}
            if har != vent:
                feil.append(f"{lab} {b['id']}: per type {har} ≠ modellrapporten {vent}")
        if sum(x["n"] for x in mine) != rm[lab]["n"]:
            feil.append(f"{lab}: objekter per type {sum(x['n'] for x in mine)} ≠ modellrapporten {rm[lab]['n']}")
        # per object: the same judgement over one object, cached by its readings
        enkel: dict[str, dict] = {}

        def status(o: dict, bdef: dict) -> str:
            bid = bdef["id"]
            kilde = "materialprodukt" if bid in ("produkt", "materiale") else (
                "mmi" if bdef.get("mmi_fase") else bid)
            ev = [e for e in o["b"] if e[0] == kilde]
            nk = json.dumps([bid, ev, o.get("mk"), o.get("s")], ensure_ascii=False)
            if nk not in enkel:
                bd_, _ = gruppe_bd([{**o, "b": ev}], blokkinfo, tuple(o["t"]) if o["t"] else None)
                b = dommer(mb, data[lab], bd_, [bdef])[bid]
                bot = next((e[2] for e in ev if e[1] == "s"), "")
                v = b["verdikt"]
                enkel[nk] = ("" if v == "ikke_konfigurert" else "gjelder_ikke" if v == "gjelder_ikke"
                             else "oppfylt" if v == "oppfylt" else "mangler" if bot == "mangler" else "avvik")
            return enkel[nk]

        for x in mine:
            for o in x["objekter"]:
                forekomster.append((x, o, {b["id"]: status(o, b) for b in kol}))
        for d in po.get("deler") or []:
            deler_alle.append((lab, d))
        rader += mine
        sammendrag.append((lab, mine, len(po.get("deler") or [])))
    if feil:
        raise SystemExit("FEIL: paritet mot modellrapporten:\n  " + "\n  ".join(feil))
    print(f"paritet mot modellrapporten: {len(labels)} modeller × {len(paritet)} krav OK")

    # ------------------------------------------------ values per type and object
    def verdier_type(x: dict, bdef: dict) -> str | None:
        bid = bdef["id"]
        b = x["blk"]["materialprodukt" if bid in ("produkt", "materiale") else bid]
        if bid in ("produkt", "materiale"):
            gren = {"produkt": "telleobjekt", "materiale": "mengdeobjekt"}[bid]
            c = b.gren_verdier.get(gren, collections.Counter())
        else:
            c = b.verdier
        deler = [f"{v} ×{n}" for v, n in c.most_common()]
        mangler = x["dom"][bid]["dekning"]["mangler"]
        if mangler:
            deler.append(f"{E('ty_uten_verdi')} ×{mangler}")
        return " · ".join(deler) or None

    def verdi_objekt(o: dict, bdef: dict) -> str | None:
        bid = bdef["id"]
        kilde = "materialprodukt" if bid in ("produkt", "materiale") else ("mmi" if bdef.get("mmi_fase") else bid)
        for e in o["b"]:
            if e[0] == kilde and e[1] == "s":
                if bid in ("produkt", "materiale") and e[5] and e[5] != {"produkt": "telleobjekt",
                                                                          "materiale": "mengdeobjekt"}[bid]:
                    return None
                return " + ".join(str(v) for v in e[4]) or None
        return None

    def klar(x: dict) -> tuple[bool, list[str]]:
        mangler = []
        for b in kol:
            if b["id"] not in T["klar"]:
                continue
            t_ = tall(x["dom"][b["id"]])
            if t_["avvik"] or t_["mangler"] or t_.get("utenfor"):
                mangler.append(b["tittel"])
        return not mangler, mangler

    li = {lab: i for i, lab in enumerate(labels)}
    rader.sort(key=lambda x: (li[x["modell"]], not x["typet"], (x["typenavn"] or x["typeklasse"]).casefold(),
                              x["typeklasse"]))
    forekomster.sort(key=lambda t: (li[t[0]["modell"]], not t[0]["typet"], t[0]["typenavn"].casefold(),
                                    t[0]["typeklasse"], t[1]["k"], t[1]["g"]))

    # ------------------------------------------------ workbook
    fyll = {s: PatternFill("solid", fgColor=MK.hex_(tok, f"--sk-{n}-flate"))
            for s, n in (("ok", "oppfylt"), ("warn", "avvik"), ("bad", "mangler"), ("na", "na"))}
    st_fyll = {"oppfylt": "ok", "avvik": "warn", "mangler": "bad", "gjelder_ikke": "na"}
    st_ord = {k: E(f"ty_st_{k}") for k in BOTTER}
    fet = Font(bold=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    def ark(navn: str, hode: list[str], band: list[tuple[str, int]] | None = None, bred: dict | None = None):
        """A sheet with the header in row 1, or with a band of section titles over it (row 1, merged
        over its columns) and the header in row 2."""
        ws = wb.create_sheet(navn)
        if band:
            ws.append([None] * len(hode))
            c0 = 1
            for tittel, n in band:
                if n:
                    ws.cell(row=1, column=c0, value=tittel).font = fet
                    if tittel and n > 1:
                        ws.merge_cells(start_row=1, start_column=c0, end_row=1, end_column=c0 + n - 1)
                c0 += n
        ws.append(hode)
        h = ws.max_row
        for c in ws[h]:
            c.font = fet
            c.alignment = Alignment(wrap_text=True, vertical="bottom")
        for i, k in enumerate(hode, 1):
            ws.column_dimensions[get_column_letter(i)].width = (bred or {}).get(k, 16)
        return ws, h

    def avslutt(ws, h: int, frys: str) -> None:
        ws.freeze_panes = frys
        if ws.max_row > h:
            ws.auto_filter.ref = f"A{h}:{get_column_letter(ws.max_column)}{ws.max_row}"

    blokk_band = [(s["tittel"], 2 * sum(1 for b in kol if b["seksjon"] == s["id"])) for s in T["seksjoner"]]
    blokk_hode = [h_ for b in kol for h_ in (b["tittel"], f"{b['tittel']} {E('ty_status')}")]

    # Sammendrag
    hode = [E(k) for k in ("ty_modell", "ty_firma", "ty_versjon", "ty_objekter", "ty_typer", "ty_utypede",
                           "ty_deler")]
    if T["klar_paa"]:
        hode += [E("ty_klar"), E("ty_krever")]
    fast = len(hode)
    ws, h = ark(E("ark_ty_sammendrag"), hode + [b["tittel"] for b in kol],
                [("", fast), (E("ty_typer_oppfylt"), len(kol))], {E("ty_modell"): 24})
    for lab, mine, n_deler in sammendrag:
        m = r.cfg[lab]
        rad = [lab, m.get("firma") or None, m.get("versjon") or None, rm[lab]["n"], len(mine),
               sum(x["n"] for x in mine if not x["typet"]), n_deler]
        if T["klar_paa"]:
            k_ = sum(1 for x in mine if klar(x)[0])
            rad += [k_, len(mine) - k_]
        for b in kol:
            gj = [x for x in mine if x["dom"][b["id"]]["verdikt"] not in ("gjelder_ikke", "ikke_konfigurert")]
            ok = sum(1 for x in gj if x["dom"][b["id"]]["verdikt"] == "oppfylt")
            rad.append(f"{ok}/{len(gj)}" if gj else MK.VERDIKT_ORD["gjelder_ikke"])
        ws.append(rad)
        for j, b in enumerate(kol):
            rb = next(x for x in rm[lab]["blokker"] if x["id"] == b["id"])
            ws.cell(row=ws.max_row, column=fast + 1 + j).fill = fyll[MK.VERDIKT_STATUS[rb["verdikt"]]]
    avslutt(ws, h, f"B{h + 1}")

    # Typer
    q_hode = [m["tittel"] for m in T["mengder"]]
    fast_t = [E(k) for k in ("ty_modell", "ty_typenavn", "ty_typeklasse", "ty_ifc_klasse", "ty_objekter")] + q_hode
    if T["klar_paa"]:
        fast_t += [T["klar_tittel"], E("ty_mangler")]
    ws, h = ark(E("ark_ty_typer"), fast_t + blokk_hode, [("", len(fast_t))] + blokk_band,
                {E("ty_modell"): 20, E("ty_typenavn"): 36, E("ty_ifc_klasse"): 22})
    for x in rader:
        kl = collections.Counter(o["k"] for o in x["objekter"])
        rad = [x["modell"], (x["typenavn"] or E("ty_uten_navn")) if x["typet"] else E("ty_utypet"),
               x["typeklasse"] or None, ", ".join(k for k, _ in kl.most_common()), x["n"]]
        for m in T["mengder"]:
            vv = [o["q"][m["id"]] for o in x["objekter"] if m["id"] in (o.get("q") or {})]
            rad.append(round(sum(vv), 3) if vv else None)
        n_q = len(T["mengder"])
        if T["klar_paa"]:
            ok, mangler = klar(x)
            rad += [E("ty_klar") if ok else E("ty_krever"), ", ".join(mangler) or None]
        for b in kol:
            rad += [verdier_type(x, b), MK.VERDIKT_ORD[x["dom"][b["id"]]["verdikt"]]]
        ws.append(rad)
        rr = ws.max_row
        for j in range(n_q):
            ws.cell(row=rr, column=6 + j).number_format = "#,##0.00"
        if T["klar_paa"]:
            ws.cell(row=rr, column=len(fast_t) - 1).fill = fyll["ok" if ok else "warn"]
        for j, b in enumerate(kol):
            ws.cell(row=rr, column=len(fast_t) + 2 * j + 2).fill = fyll[MK.VERDIKT_STATUS[x["dom"][b["id"]]["verdikt"]]]
    avslutt(ws, h, f"C{h + 1}")

    # Forekomster: every counted object, then the parts of an assembly with their owner
    fast_f = [E(k) for k in ("ty_modell", "ty_typenavn", "ty_ifc_klasse", "ty_globalid", "ty_type_globalid",
                             "ty_etasje", "ty_eier_guid", "ty_eier_type")]
    # one column per requirement: the value, filled by its status; no value is «missing» in grey
    blokk_hode_f = [b["tittel"] for b in kol]
    blokk_band_f = [(s["tittel"], sum(1 for b in kol if b["seksjon"] == s["id"])) for s in T["seksjoner"]]
    ws, h = ark(E("ark_ty_forekomster"), fast_f + blokk_hode_f, [("", len(fast_f))] + blokk_band_f,
                {E("ty_typenavn"): 36, E("ty_globalid"): 24, E("ty_type_globalid"): 24, E("ty_eier_guid"): 24})
    for x, o, st in forekomster:
        et = o["s"][0] if o.get("s") else (f"({o.get('h')})" if o.get("h") else None)
        rad = [x["modell"], (x["typenavn"] or E("ty_uten_navn")) if x["typet"] else None, o["k"], o["g"],
               o["tg"] or None, et, None, None]
        celler = []
        for b in kol:
            sb, v = st[b["id"]], verdi_objekt(o, b)
            if not sb:
                celler.append((None, None))
            elif sb == "gjelder_ikke":
                celler.append((st_ord[sb], "na"))
            elif sb == "mangler" or not v:
                celler.append((st_ord["mangler"], "na"))
            else:
                # A value that fails is red, not the amber «avvik» of the verdicts (edkjo 2026-10-30).
                celler.append((v, "ok" if sb == "oppfylt" else "bad"))
        rad += [c[0] for c in celler]
        ws.append(rad)
        rr = ws.max_row
        for j, (_, f_) in enumerate(celler):
            if f_:
                ws.cell(row=rr, column=len(fast_f) + 1 + j).fill = fyll[f_]
    for lab, d in deler_alle:
        ws.append([lab, d["t"], d["k"], d["g"], d.get("tg") or None, None, d["eier"],
                   d["eier_t"] if d["eier_t"] is not None else d["eier_k"]] + [None] * len(blokk_hode_f))
    avslutt(ws, h, f"E{h + 1}")

    # Materialer: every material name the counted objects carry, before any join
    mat = collections.defaultdict(lambda: [set(), set(), 0])
    for x in rader:
        for o in x["objekter"]:
            for v in dict.fromkeys(o.get("mn") or []) or [E("ty_uten_verdi")]:
                e_ = mat[v]
                e_[0].add(x["modell"])
                e_[1].add((x["modell"], x["typeklasse"], x["typenavn"], x["typet"] or o["k"]))
                e_[2] += 1
    ws, h = ark(E("ark_ty_materialer"), [E(k) for k in ("ty_materiale", "ty_modeller", "ty_typer", "ty_objekter")],
                bred={E("ty_materiale"): 40, E("ty_modeller"): 30})
    for v, (mods, typer, n) in sorted(mat.items(), key=lambda kv: (-kv[1][2], kv[0].casefold())):
        ws.append([v, ", ".join(lab for lab in labels if lab in mods), len(typer), n])
    avslutt(ws, h, f"B{h + 1}")

    # Materiallag: the layers as the model states them, the type's association first, else the
    # objects' own (the most common)
    typemat = {}
    for lab in labels:
        bd0 = json.loads((r.blokk_sti / "prosjekt" / f"{lab}.json").read_text(encoding="utf-8"))
        typemat[lab] = bd0["per_objekt"].get("typemateriale") or {}
    ml_rader = []
    for x in rader:
        tm = typemat[x["modell"]].get(f"{x['typeklasse']}|{x['typenavn']}") if x["typet"] else None
        if tm:
            kilde, lag = E("ty_kilde_type"), tm
        else:
            c = collections.Counter(json.dumps(o["ml"], ensure_ascii=False) for o in x["objekter"] if o.get("ml"))
            kilde, lag = (E("ty_kilde_forekomst"), json.loads(c.most_common(1)[0][0])) if c else (None, [])
        ml_rader.append((x, kilde, lag))
    bredest = max([1] + [len(l_) for _, _, l_ in ml_rader])
    hode = [E(k) for k in ("ty_modell", "ty_typenavn", "ty_typeklasse", "ty_mat_kilde", "ty_antall_lag",
                           "ty_mengde_per_lag")]
    for i in range(1, bredest + 1):
        hode += [f"{E('ty_lag')} {i} {E('ty_lag_materiale')}", f"{E('ty_lag')} {i} {E('ty_lag_tykkelse')}",
                 f"{E('ty_lag')} {i} {E('ty_lag_andel')}"]
    ws, h = ark(E("ark_ty_materiallag"), hode, bred={E("ty_typenavn"): 36})
    for x, kilde, lag in ml_rader:
        arter = list(dict.fromkeys(E(f"ty_lagart_{a}") for _, _, a in lag))
        rad = [x["modell"], (x["typenavn"] or E("ty_uten_navn")) if x["typet"] else E("ty_utypet"),
               x["typeklasse"] or None, kilde, len(lag), ", ".join(arter) or None]
        for navn_, mengde, art in lag:
            rad += [navn_ or None, mengde if art == "tykkelse" else None, mengde if art == "andel" else None]
        ws.append(rad + [None] * (len(hode) - len(rad)))
        for i in range(len(lag)):
            ws.cell(row=ws.max_row, column=8 + 3 * i).number_format = "0.0"
            ws.cell(row=ws.max_row, column=9 + 3 * i).number_format = "0.0%"
    avslutt(ws, h, f"C{h + 1}")

    # Kilder
    ws, h = ark(E("ark_ty_kilder"), [E(k) for k in ("ty_modell", "ty_fil", "ty_sha16", "ty_versjon",
                                                     "ty_lastet_opp", "ty_firma", "ty_objekter")],
                bred={E("ty_modell"): 30, E("ty_fil"): 30})
    for lab in labels:
        m = r.cfg[lab]
        ws.append([lab, m["fil"], m["sha16"], m.get("versjon") or None, m.get("lastet_opp") or None,
                   m.get("firma") or None, rm[lab]["n"]])
    ws.append([])
    konfig_sha = hashlib.sha256(prosjekt.read_bytes()).hexdigest()[:16]
    fakta = [(E("ty_runde"), f"{r.dato} ({runde.name})"), (E("ty_bygget"), dt.date.today().isoformat()),
             (E("ty_konfig"), f"{prosjekt.name} · sha16 {konfig_sha}"),
             (E("ty_paritet"), f"{E('ty_paritet_tekst')} ({len(labels)} × {len(paritet)})")]
    for lab, grunn in T["utelatt"].items():
        if lab in hoved:
            fakta.append((E("ty_utelatt_modell"), f"{lab}: {grunn}"))
    for k, v in fakta:
        ws.append([k, v])
        ws.cell(row=ws.max_row, column=1).font = fet
    avslutt(ws, h, f"B{h + 1}")

    ut.mkdir(parents=True, exist_ok=True)
    sti = ut / T["filnavn"]
    wb.save(sti)
    moj = MK.sjekk_mojibake([sti])
    if moj:
        print("\n".join(f"FEIL: mojibake: {x}" for x in moj))
        return 1
    print(f"{sti.name}: {len(rader)} typer, {len(forekomster)} objekter, {len(deler_alle)} deler")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description="The type report of one round, from its measurements.")
    stier.argumenter(ap)
    ap.add_argument("--ifc", type=Path, metavar="DIR", help="IFC folder of the round (replaces the round's `kilde`)")
    a = ap.parse_args()
    return bygg(a.prosjekt, a.runde, a.cache, a.ut, a.ifc)


if __name__ == "__main__":
    raise SystemExit(main())
